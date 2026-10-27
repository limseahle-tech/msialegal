#!/usr/bin/env python3
"""Build a Malaysian court submission (.docx) in Caroline's house format.

Usage:
    python3 build_submission.py draft.md -o "Pf Reply Submissions - v1.docx"

Format enforced: Arial 12, line spacing 1.5, 0 pt before/after, an empty
paragraph between paragraphs, indentation Left 0" / Right 0" / Hanging 0.5"
for every body paragraph, and Word automatic multi-level numbering
(1. -> (a) -> (i)).

INPUT FILE
----------
Front matter between two '---' lines, one 'key: value' per line.
Use ' | ' to split a value into several lines.

    court:      DALAM MAHKAMAH TINGGI MALAYA DI KUALA LUMPUR | DALAM WILAYAH PERSEKUTUAN KUALA LUMPUR, MALAYSIA | (BAHAGIAN DAGANG)
    suit:       GUAMAN SIVIL NO.: WA-22IP-21-02/2026
    plaintiff:  NAME | (NO. SYARIKAT: 201301044553 [1074376-V])     (repeatable)
    plaintiff_role: ...PLAINTIF
    defendant:  NAME | (NO. K/P: 620909-02-5160)                     (repeatable)
    defendant_role: ...DEFENDAN-DEFENDAN
    title:      PLAINTIFF'S SUBMISSIONS IN REPLY
    subtitle:   (in opposing Enclosure 16)
    solicitors_left:  **Solicitors for the Plaintiff** | **Messrs Lum Kok Kiong & Co.** | K-6-6, Solaris Mont' Kiara, | ... | [*Ref.: L/LKK/...*]
    solicitors_right: (optional, same format)
    date:       Dated this 12th of June, 2026.
    signatory:  Solicitors for the Plaintiff | Messrs Lum Kok Kiong & Co
    filed_by:   This **PLAINTIFF'S SUBMISSIONS IN REPLY** *(in opposing Enclosure 16)* is filed by ...
    contact:    [Tel: 03-6420 7841] [Fax: 03-6420 7842] | [Email: ...] [Ref.: ...]
    opening:    (optional) replaces "May it please this Honourable Court," e.g.
                *Dengan Izin Mahkamah yang Mulia ini,* | Kami memohon kebenaran untuk berhujah dalam Bahasa Inggeris.
    sublevel:   (optional) letter (default: 1. -> (a) -> (i)) or decimal (1. -> 1.1 -> (a))

Body (after the front matter), one paragraph per line:

    # INTRODUCTION            top-level heading: bold, underlined
    ## NO COMMON QUESTION     argument sub-heading: bold, not underlined
    ### Legal principles      third level: bold italic
    1. text                   numbered paragraph (number is automatic)
       (a) text               sub-paragraph, 3+ leading spaces (letter automatic)
          (i) text            sub-sub-paragraph, 6+ leading spaces
    > text                    quotation (italic, unnumbered)
    text                      plain unnumbered paragraph
    | a | b |                 table row (first row bold; |---| rows ignored)
    \\toc                     table of contents field (update in Word: F9)

Inline: **bold**  *italic*  __underline__  ^st^ superscript  ==XX== yellow highlight (placeholders)
        [[Case Name [2016] 5 MLJ 417 (FC)]]  -> bold + italic + underline
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Arial"
SIZE = Pt(12)
HANG = Inches(0.5)
NUM_ID = 901  # high id to avoid clashing with template numbering

INLINE = re.compile(r"(\[\[.*?\]\](?!\])|==.+?==|\*\*|__|\*|\^[^^\s]+\^)")


# ---------------------------------------------------------------- inline
def add_runs(par, text, bold=False, italic=False, underline=False):
    state = {"b": bold, "i": italic, "u": underline}
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok == "**":
            state["b"] = not state["b"]
        elif tok == "*":
            state["i"] = not state["i"]
        elif tok == "__":
            state["u"] = not state["u"]
        elif tok.startswith("[[") and tok.endswith("]]"):
            _run(par, tok[2:-2], True, True, True)
        elif tok.startswith("==") and tok.endswith("==") and len(tok) > 4:
            r = _run(par, tok[2:-2], state["b"], state["i"], state["u"])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif tok.startswith("^") and tok.endswith("^") and len(tok) > 2:
            r = _run(par, tok[1:-1], state["b"], state["i"], state["u"])
            r.font.superscript = True
        else:
            _run(par, tok, state["b"], state["i"], state["u"])


def _run(par, text, b, i, u):
    r = par.add_run(text)
    r.bold, r.italic, r.underline = b, i, u
    _font(r)
    return r


def _font(run):
    run.font.name = FONT
    run.font.size = SIZE
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), FONT)


# ---------------------------------------------------------------- format
def fmt(par, hanging=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = par.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.right_indent = Inches(0)
    if hanging:
        pf.left_indent = HANG
        pf.first_line_indent = -HANG
    else:
        pf.left_indent = Inches(0)
        pf.first_line_indent = Inches(0)
    par.alignment = align
    return par


def blank(doc, keep=False):
    p = fmt(doc.add_paragraph())
    p.paragraph_format.keep_with_next = keep


def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = SIZE
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)


def setup_numbering(doc, sublevel="letter"):
    numbering = doc.part.numbering_part.element
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), str(NUM_ID))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "multilevel")
    absn.append(mlt)
    if sublevel == "decimal":
        levels = [("decimal", "%1."), ("decimal", "%1.%2"), ("lowerLetter", "(%3)")]
    else:
        levels = [("decimal", "%1."), ("lowerLetter", "(%2)"), ("lowerRoman", "(%3)")]
    for ilvl, (nfmt, text) in enumerate(levels):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", nfmt),
                         ("w:lvlText", text), ("w:lvlJc", "left")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        if ilvl:
            e = OxmlElement("w:lvlRestart")
            e.set(qn("w:val"), str(ilvl))
            lvl.insert(1, e)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:right"), "0")
        ind.set(qn("w:hanging"), "720")
        ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        f = OxmlElement("w:rFonts")
        for k in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(k), FONT)
        rpr.append(f)
        lvl.append(rpr)
        absn.append(lvl)
    # abstractNum elements must precede num elements
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(absn)
    else:
        numbering.append(absn)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(NUM_ID))
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), str(NUM_ID))
    num.append(a)
    numbering.append(num)


def numbered(doc, text, level):
    p = fmt(doc.add_paragraph())
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), str(level))
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(NUM_ID))
    numpr.append(il)
    numpr.append(ni)
    ppr.insert(0, numpr)
    add_runs(p, text)
    return p


# ---------------------------------------------------------------- tables
def _borders(table, on):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if on else "nil")
        if on:
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), "000000")
        b.append(e)
    tblpr.append(b)


def cell_lines(cell, lines, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for line in lines:
        p = fmt(cell.add_paragraph(), hanging=False, align=align)
        add_runs(p, line, bold=bold)


def split(v):
    return [x.strip() for x in v.split(" | ")]


# ---------------------------------------------------------------- front/back
def front(doc, m):
    for line in split(m.get("court", "")):
        if line:
            add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), line, bold=True)
    blank(doc)
    if m.get("suit"):
        add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), m["suit"], bold=True, underline=True)
        blank(doc)

    def party_table(names, role, numbered_rows):
        t = doc.add_table(rows=0, cols=3 if numbered_rows else 2)
        t.autofit = False
        _borders(t, False)
        widths = [Inches(0.5), Inches(3.8), Inches(2.2)] if numbered_rows else [Inches(4.3), Inches(2.2)]
        for idx, nm in enumerate(names, 1):
            row = t.add_row().cells
            if numbered_rows:
                cell_lines(row[0], [f"{idx}."], bold=True)
                cell_lines(row[1], split(nm), bold=True)
                cell_lines(row[2], [role if idx == len(names) else ""], bold=True)
            else:
                cell_lines(row[0], split(nm), bold=True)
                cell_lines(row[1], [role if idx == len(names) else ""], bold=True)
            for c, w in zip(row, widths):
                c.width = w
        for col, w in zip(t.columns, widths):
            col.width = w
        blank(doc)

    if m.get("plaintiff"):
        add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), "ANTARA", bold=True)
        blank(doc)
        pls = m["plaintiff"]
        party_table(pls, m.get("plaintiff_role", "...PLAINTIF"), len(pls) > 1)
    if m.get("defendant"):
        add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), "DAN", bold=True)
        blank(doc)
        dfs = m["defendant"]
        party_table(dfs, m.get("defendant_role", "...DEFENDAN"), len(dfs) > 1)

    if m.get("title"):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, True)
        lines = [f"**{m['title']}**"] + ([f"*{m['subtitle']}*"] if m.get("subtitle") else [])
        cell_lines(t.rows[0].cells[0], lines, align=WD_ALIGN_PARAGRAPH.CENTER)
        blank(doc)

    if m.get("solicitors_left"):
        cols = 2 if m.get("solicitors_right") else 1
        t = doc.add_table(rows=1, cols=cols)
        _borders(t, True)
        cell_lines(t.rows[0].cells[0], split(m["solicitors_left"]))
        if cols == 2:
            cell_lines(t.rows[0].cells[1], split(m["solicitors_right"]))
        doc.add_page_break()

    if m.get("title"):
        add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), m["title"], bold=True, underline=True)
        if m.get("subtitle"):
            add_runs(fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.CENTER), m["subtitle"], italic=True)
        blank(doc)
    for line in split(m.get("opening", "May it please this Honourable Court,")):
        add_runs(fmt(doc.add_paragraph(), False), line)
    blank(doc)


def back(doc, m):
    if m.get("date"):
        p = fmt(doc.add_paragraph(), False)
        p.paragraph_format.keep_with_next = True
        add_runs(p, m["date"])
        blank(doc, keep=True)
    if m.get("signatory"):
        t = doc.add_table(rows=1, cols=2)
        _borders(t, False)
        cell_lines(t.rows[0].cells[0], [""])
        cell_lines(t.rows[0].cells[1], ["", "", "." * 45] + split(m["signatory"]),
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in t.rows[0].cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
        blank(doc)
    if m.get("filed_by"):
        add_runs(fmt(doc.add_paragraph(), False), m["filed_by"])
        blank(doc)
    for line in split(m.get("contact", "")):
        if line:
            add_runs(fmt(doc.add_paragraph(), False), line)


def toc(doc):
    p = fmt(doc.add_paragraph(), False)
    r = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None),
                       (None, None), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "[Right-click and Update Field to build the Table of Contents]"
            r._r.append(t)
    doc.add_page_break()


# ---------------------------------------------------------------- body
def outline(par, level):
    """Mark a heading for the Table of Contents without using Heading styles."""
    e = OxmlElement("w:outlineLvl")
    e.set(qn("w:val"), str(level))
    par._p.get_or_add_pPr().append(e)


def body(doc, lines):
    rows = []

    def flush_table():
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncols)
        _borders(t, True)
        for i, r in enumerate(rows):
            cells = t.add_row().cells
            for j in range(ncols):
                txt = r[j] if j < len(r) else ""
                cell_lines(cells[j], txt.split("<br>"), bold=(i == 0))
        rows.clear()
        blank(doc)

    for raw in lines:
        line = raw.rstrip("\n")
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                rows.append(cells)
            continue
        flush_table()
        if not stripped:
            continue
        indent = len(line) - len(line.lstrip(" "))
        if stripped == "\\toc":
            toc(doc)
        elif stripped.startswith("### "):
            p = fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.LEFT)
            add_runs(p, stripped[4:], bold=True, italic=True)
            outline(p, 2)
            p.paragraph_format.keep_with_next = True
            blank(doc, keep=True)
        elif stripped.startswith("## "):
            p = fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.LEFT)
            add_runs(p, stripped[3:], bold=True)
            outline(p, 1)
            p.paragraph_format.keep_with_next = True
            blank(doc, keep=True)
        elif stripped.startswith("# "):
            p = fmt(doc.add_paragraph(), False, WD_ALIGN_PARAGRAPH.LEFT)
            add_runs(p, stripped[2:], bold=True, underline=True)
            outline(p, 0)
            p.paragraph_format.keep_with_next = True
            blank(doc, keep=True)
        elif stripped.startswith(">"):
            add_runs(fmt(doc.add_paragraph()), stripped.lstrip("> ").rstrip(), italic=True)
            blank(doc)
        elif re.match(r"^\d+\.\s", stripped) and indent < 3:
            numbered(doc, re.sub(r"^\d+\.\s+", "", stripped), 0)
            blank(doc)
        elif re.match(r"^(\([a-z]{1,4}\)|[a-z]{1,4}\.|\d+\.\d+)\s", stripped) and indent >= 3:
            level = 2 if indent >= 6 else 1
            numbered(doc, re.sub(r"^(\([a-z]{1,4}\)|[a-z]{1,4}\.|\d+\.\d+)\s+", "", stripped), level)
            blank(doc)
        else:
            add_runs(fmt(doc.add_paragraph()), stripped)
            blank(doc)
    flush_table()


def parse(path):
    with open(path, encoding="utf-8") as fh:
        lines = fh.readlines()
    meta, start = {}, 0
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                start = i + 1
                break
            if ":" in lines[i]:
                k, v = lines[i].split(":", 1)
                k, v = k.strip(), v.strip()
                if k in ("plaintiff", "defendant"):
                    meta.setdefault(k, []).append(v)
                else:
                    meta[k] = v
    return meta, lines[start:]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("input")
    ap.add_argument("-o", "--output", required=True)
    a = ap.parse_args()
    meta, lines = parse(a.input)
    doc = Document()
    setup_styles(doc)
    setup_numbering(doc, meta.get("sublevel", "letter"))
    front(doc, meta)
    body(doc, lines)
    back(doc, meta)
    doc.save(a.output)
    print(f"Saved {a.output}")


if __name__ == "__main__":
    sys.exit(main())
